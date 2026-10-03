from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "docs/cloud-devops-feature-list.docx"


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
TEXT = "1F2937"
MUTED = "4B5563"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=TEXT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_table_borders(table, color="C9D3E0"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor.from_string(BLUE if level <= 2 else DARK_BLUE)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor.from_string(TEXT)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(TEXT)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(TEXT)
    return p


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_widths(table, [6.5])
    set_table_borders(table, "D6DEE8")
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(text)
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor.from_string(TEXT)
    doc.add_paragraph()


def add_feature_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    set_table_widths(table, [1.55, 2.35, 1.1, 1.5])
    set_table_borders(table)
    headers = ["Feature", "What It Adds", "Difficulty", "Resume Value"]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_text(cell, header, bold=True, color=DARK_BLUE)
    for feature, adds, difficulty, value in rows:
        cells = table.add_row().cells
        values = [feature, adds, difficulty, value]
        for idx, value_text in enumerate(values):
            set_cell_text(cells[idx], value_text)
    doc.add_paragraph()


def add_feature_section(doc, title, explanation, implementation, course_value, resume_line):
    add_heading(doc, title, 2)
    add_body(doc, explanation)
    add_bullet(doc, f"Implementation: {implementation}")
    add_bullet(doc, f"Cloud/DevOps value: {course_value}")
    add_bullet(doc, f"Resume line: {resume_line}")


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def main():
    Path(OUTPUT).parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Cloud Computing and DevOps Feature Plan")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    sub = subtitle.add_run("For the React + FastAPI + PostgreSQL Hotel Management System")
    sub.font.name = "Calibri"
    sub.font.size = Pt(11)
    sub.font.color.rgb = RGBColor.from_string(MUTED)

    add_callout(
        doc,
        "Recommended project positioning",
        "Present the system as a cloud-native hotel management platform with containerized services, automated CI/CD, managed database deployment, cloud storage, and monitoring. This makes it suitable for a Cloud Computing and DevOps course as well as a resume portfolio project.",
    )

    add_heading(doc, "Feature Overview", 1)
    add_body(
        doc,
        "The current application already has a strong base: a React frontend, a FastAPI backend, PostgreSQL, authentication, and role-based hotel workflows. The features below add the cloud and DevOps layer around that application so it can be demonstrated as a deployable, maintainable system.",
    )
    add_feature_table(
        doc,
        [
            ("Dockerization", "Package frontend, backend, and database into repeatable containers.", "Easy", "High"),
            ("Docker Compose", "Run the complete stack locally with one command.", "Easy", "Very High"),
            ("Automated DB Init", "Create schema and seed demo data automatically in PostgreSQL.", "Easy", "High"),
            ("Cloud PostgreSQL", "Use a managed Postgres service such as Neon, Supabase, Railway, or AWS RDS.", "Medium", "High"),
            ("Cloud Deployment", "Deploy backend API and frontend UI to public cloud platforms.", "Medium", "Very High"),
            ("GitHub Actions CI", "Run backend checks and frontend builds automatically on push.", "Medium", "Very High"),
            ("CI/CD Deployment", "Automatically deploy successful main-branch builds.", "Medium", "Very High"),
            ("Cloud Storage", "Upload room images, ID proof files, maintenance photos, or invoices.", "Medium", "Very High"),
            ("Health Checks", "Expose backend and database readiness endpoints for cloud monitoring.", "Easy", "High"),
            ("Logging/Monitoring", "Capture application logs and optionally integrate Sentry or uptime checks.", "Medium", "High"),
        ],
    )

    add_heading(doc, "Detailed Feature List", 1)

    add_feature_section(
        doc,
        "1. Dockerization",
        "Dockerization means packaging each part of the system into its own container. The backend would run inside a Python container, the frontend inside a Node/Nginx container, and PostgreSQL inside an official database container.",
        "Add backend/Dockerfile, frontend/Dockerfile, and .dockerignore files. The backend image installs Python dependencies and runs Uvicorn. The frontend image builds the Vite app and serves the production files.",
        "Shows containerization, consistent environments, and modern application packaging.",
        "Containerized a full-stack hotel management system using Docker.",
    )

    add_feature_section(
        doc,
        "2. Docker Compose Full-Stack Setup",
        "Docker Compose will define the frontend, backend, and PostgreSQL services in one YAML file. This allows the entire project to start with a single command instead of manually opening multiple terminals and installing PostgreSQL locally.",
        "Add docker-compose.yml with services for frontend, backend, postgres, volumes, ports, and environment variables.",
        "Demonstrates multi-container orchestration and service networking.",
        "Built a Docker Compose environment for frontend, API, and database services.",
    )

    add_feature_section(
        doc,
        "3. Automated Database Initialization",
        "The PostgreSQL container can automatically create tables and insert demo data when it starts for the first time. This removes the need to manually run SQL files in pgAdmin for local development.",
        "Mount database/complete_schema.sql and database/seed_demo.sql into the PostgreSQL initialization directory in Docker.",
        "Shows repeatable database provisioning and reproducible setup.",
        "Automated PostgreSQL schema creation and seed data loading for local deployments.",
    )

    add_feature_section(
        doc,
        "4. Managed Cloud PostgreSQL",
        "A managed cloud database moves production data away from a local computer. The backend connects to a hosted PostgreSQL database through DATABASE_URL.",
        "Use Neon, Supabase, Railway PostgreSQL, or AWS RDS. Update backend environment variables to use the hosted connection string.",
        "Demonstrates database-as-a-service and production-style configuration.",
        "Integrated managed PostgreSQL for cloud-hosted application data.",
    )

    add_feature_section(
        doc,
        "5. Backend Cloud Deployment",
        "The FastAPI backend should be deployed to a cloud platform so the API is available through a public HTTPS URL.",
        "Deploy the backend Docker image or Python service to Render, Railway, Fly.io, AWS Elastic Beanstalk, or AWS ECS.",
        "Shows API deployment, environment variables, cloud logs, and production backend hosting.",
        "Deployed a FastAPI backend to cloud hosting with production environment variables.",
    )

    add_feature_section(
        doc,
        "6. Frontend Cloud Deployment",
        "The React frontend should be deployed separately so users can access the hotel management dashboard from a browser without running local tools.",
        "Deploy the Vite React frontend to Vercel, Netlify, or Cloudflare Pages. Configure VITE_API_URL to point to the cloud backend.",
        "Demonstrates static frontend hosting and API integration.",
        "Deployed a React frontend to cloud hosting and connected it to a live FastAPI API.",
    )

    add_feature_section(
        doc,
        "7. GitHub Actions CI Pipeline",
        "Continuous Integration automatically checks whether the project still builds whenever code is pushed to GitHub.",
        "Add .github/workflows/ci.yml to install backend dependencies, compile/check Python, install frontend dependencies, and build the React app.",
        "Shows automated validation, team collaboration safety, and DevOps workflow knowledge.",
        "Implemented GitHub Actions CI for automated backend validation and frontend builds.",
    )

    add_feature_section(
        doc,
        "8. CI/CD Deployment Pipeline",
        "Continuous Deployment extends CI by deploying the project automatically after successful checks on the main branch.",
        "Use GitHub Actions with deployment hooks or platform integrations for Vercel, Render, Railway, or similar services.",
        "Shows automated release management and production delivery.",
        "Configured CI/CD deployment from GitHub to cloud-hosted frontend and backend services.",
    )

    add_feature_section(
        doc,
        "9. Environment and Secret Management",
        "The project should separate local development configuration from production configuration. Secrets such as database passwords and JWT keys should never be committed to GitHub.",
        "Use .env.example files locally and store production values in GitHub Secrets or cloud platform environment variables.",
        "Shows secure configuration practices and production readiness.",
        "Implemented environment-based configuration and secret management for cloud deployment.",
    )

    add_feature_section(
        doc,
        "10. Cloud File Storage",
        "A hotel system naturally needs images and documents. Cloud storage can be used for room photos, customer ID proof uploads, invoice PDFs, and maintenance issue photos.",
        "Use Cloudinary, Supabase Storage, or AWS S3. Backend uploads files and stores the returned file URL in PostgreSQL.",
        "Adds a real cloud service integration beyond simple hosting.",
        "Integrated cloud storage for hotel images and customer document uploads.",
    )

    add_feature_section(
        doc,
        "11. Email Notifications",
        "The system can send emails when a booking is created, approved, rejected, checked in, checked out, or when an invoice/payment is generated.",
        "Use Resend, SendGrid, or AWS SES from the FastAPI backend. Trigger emails from booking and billing workflows.",
        "Shows third-party cloud service integration and event-based application design.",
        "Added automated email notifications for booking and billing workflows.",
    )

    add_feature_section(
        doc,
        "12. Health Checks",
        "Cloud platforms need health checks to know whether the application is running correctly. The existing /api/health endpoint can be improved to also verify database connectivity.",
        "Add /api/health/db or extend /api/health to return backend status, database status, app version, and environment.",
        "Shows operational readiness and deployment monitoring basics.",
        "Implemented backend and database health checks for cloud monitoring.",
    )

    add_feature_section(
        doc,
        "13. Logging and Monitoring",
        "When deployed to the cloud, debugging is done through logs and monitoring dashboards. The backend should log important events and errors clearly.",
        "Add structured Python logging for authentication, bookings, payments, and failures. Optionally integrate Sentry or uptime monitoring.",
        "Shows observability, production debugging, and maintenance awareness.",
        "Added structured logging and monitoring hooks for cloud-deployed services.",
    )

    add_feature_section(
        doc,
        "14. Backup and Migration Strategy",
        "A cloud database should have a simple backup and migration story. This makes the system more credible as a production-style application.",
        "Document database backup commands and optionally add migration tooling such as Alembic later.",
        "Shows data safety and lifecycle management.",
        "Documented PostgreSQL backup and migration strategy for application reliability.",
    )

    add_feature_section(
        doc,
        "15. Architecture and Deployment Documentation",
        "A cloud course project should clearly explain how the system is built, deployed, and operated.",
        "Add docs/architecture.md, docs/deployment.md, docs/devops-pipeline.md, screenshots, and a README architecture diagram.",
        "Helps evaluators understand the cloud design and makes the GitHub repository stronger.",
        "Documented cloud architecture, deployment process, and CI/CD pipeline.",
    )

    add_heading(doc, "Recommended Implementation Roadmap", 1)
    for item in [
        "Phase 1: Add Dockerfiles and docker-compose.yml for local containerized execution.",
        "Phase 2: Automate PostgreSQL schema and seed data initialization inside Docker.",
        "Phase 3: Add GitHub Actions CI for backend validation and frontend build checks.",
        "Phase 4: Deploy PostgreSQL to Neon or Supabase and connect the backend using DATABASE_URL.",
        "Phase 5: Deploy backend to Render or Railway and frontend to Vercel.",
        "Phase 6: Add cloud file storage for room images or customer ID proof uploads.",
        "Phase 7: Add health checks, logging, deployment screenshots, and final documentation.",
    ]:
        add_numbered(doc, item)

    add_heading(doc, "Recommended Final Feature Set", 1)
    add_body(
        doc,
        "For the best balance of course value, implementation effort, and resume strength, the project should include Docker, Docker Compose, automated PostgreSQL initialization, GitHub Actions CI, cloud PostgreSQL, deployed frontend, deployed backend, cloud file storage, health checks, and deployment documentation.",
    )

    add_heading(doc, "Resume Positioning", 1)
    add_callout(
        doc,
        "Suggested resume title",
        "Cloud-Native Hotel Management System with Dockerized Services, CI/CD Automation, Managed PostgreSQL, and Cloud Storage",
    )
    add_body(
        doc,
        "Suggested resume bullet: Built and deployed a cloud-native hotel management platform using React, FastAPI, PostgreSQL, Docker, GitHub Actions, and managed cloud services, featuring role-based workflows, automated database provisioning, CI/CD validation, cloud file storage, and production health checks.",
    )

    doc.core_properties.title = "Cloud Computing and DevOps Feature Plan"
    doc.core_properties.subject = "Feature list for converting a hotel management system into a cloud and DevOps course project"
    doc.core_properties.author = "Codex"
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
