from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


TEMPLATE = Path(r"C:\Users\nsleh\Downloads\CCD_Miniproject_Synopsys_template.docx")
OUTPUT = Path("docs/cloud-native-hotel-management-synopsis.docx")


def clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def format_runs(paragraph, size=11, bold=False, color="000000"):
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        run.font.size = Pt(size)
        run.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style="Normal", size=11, bold=False, color="000000", align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if text:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(6)
    format_runs(p, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text):
    p = add_para(doc, text, style="Heading 2", size=13, bold=True, color="1F4E79")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_label_para(doc, label, value):
    p = doc.add_paragraph(style="List Paragraph")
    p.paragraph_format.space_after = Pt(3)
    r1 = p.add_run(label)
    r1.bold = True
    r1.font.name = "Aptos"
    r1._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r1._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r1.font.size = Pt(11)
    r2 = p.add_run(value)
    r2.font.name = "Aptos"
    r2._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r2._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r2.font.size = Pt(11)
    return p


def add_module(doc, title, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(6)
    r1 = p.add_run(title + ":\n")
    r1.bold = True
    r2 = p.add_run(text)
    for run in p.runs:
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        run.font.size = Pt(11)
    return p


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.font.name = "Aptos"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r.font.size = Pt(10)
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(TEMPLATE)
    clear_body(doc)

    add_para(
        doc,
        "Cloud-Native Hotel Management System with DevOps CI/CD Pipeline",
        style="Normal",
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_heading(doc, "Student Details")
    add_label_para(doc, "Student Name(s): ", "<<Enter Student Name(s) and Registration Number(s)>>")
    add_label_para(doc, "Semester: ", "<<Enter Semester>>")
    add_label_para(doc, "Section: ", "<<Enter Section>>")
    add_label_para(doc, "Batch: ", "<<Enter Batch>>")
    add_label_para(doc, "Team : ", "<<Ask Faculty>>")
    add_label_para(doc, "Academic Year: ", "2026-27")
    add_label_para(doc, "Department: ", "School of Computer Engineering")

    add_heading(doc, "Supervisor Details")
    add_label_para(doc, "Supervisor Name: ", "<<Enter Supervisor Name>>")
    add_label_para(doc, "Department: ", "School of Computer Engineering")

    add_heading(doc, "Synopsis / Abstract")
    add_para(
        doc,
        "The proposed project is a Cloud-Native Hotel Management System developed as a full-stack web application using React for the frontend, FastAPI for the backend, and PostgreSQL for persistent data storage. The application is intended to digitalize the core operations of a hotel, including user authentication, role-based dashboards, room and room-type management, room availability tracking, customer booking, booking approval, check-in and check-out workflow, invoice generation, payment recording, hotel service usage, staff and maintenance management, and reporting.",
    )
    add_para(
        doc,
        "Along with the application development, the project focuses on converting the system into a cloud computing and DevOps based solution. The frontend, backend, and database will be containerized using Docker and executed together using Docker Compose for local development. For cloud deployment, the frontend can be hosted on Vercel or Netlify, the FastAPI backend can be deployed on Render, Railway, Fly.io, or AWS, and PostgreSQL can be moved to a managed cloud database such as Neon, Supabase, Railway PostgreSQL, or AWS RDS. GitHub Actions will be used for CI/CD automation, while environment variables, secrets management, health checks, logging, and cloud object storage will make the system closer to a real production-ready cloud application.",
    )

    add_heading(doc, "Problem Statement")
    add_para(
        doc,
        "Hotels require an efficient system to manage rooms, customers, bookings, payments, staff activities, and reports. In many small or academic systems, these operations are either handled manually, through disconnected tools, or through local desktop applications. Such approaches make it difficult to provide real-time access to multiple users, maintain consistent records, manage booking status accurately, generate invoices efficiently, and monitor hotel operations from a centralized interface.",
    )
    add_para(
        doc,
        "From a technical point of view, locally installed applications are also difficult to deploy, scale, share among team members, and maintain in a production-like environment. They often depend on manual database setup, machine-specific configuration, and manual deployment steps. This project addresses both problems by building a complete hotel management web application and then applying cloud computing and DevOps practices so that the system becomes portable, deployable, secure, automated, and easier to maintain.",
    )

    add_heading(doc, "Proposed Solution")
    add_para(
        doc,
        "The proposed solution is to develop a web-based hotel management platform and deploy it using a cloud-native architecture. The system will be divided into frontend, backend, database, cloud infrastructure, and DevOps automation layers. This layered design separates user interaction, business logic, data storage, deployment, and operational concerns, making the project suitable for both application development and cloud computing evaluation.",
    )
    add_module(
        doc,
        "Frontend Application Layer",
        "The frontend will be built using React with Vite. It will provide a responsive dashboard interface for different user roles such as admin, manager, staff, and customer. The UI will include login and registration screens, room browsing, room management, booking views, invoice and payment views, service usage screens, reports, and operational dashboards.",
    )
    add_module(
        doc,
        "Backend API Layer",
        "The backend will be developed using FastAPI in Python. It will expose REST APIs for authentication, authorization, room management, booking lifecycle operations, invoice generation, payment recording, service usage, staff and maintenance management, and reports. JWT authentication will be used to secure protected routes.",
    )
    add_module(
        doc,
        "Database Layer",
        "PostgreSQL will be used to store structured hotel data. The database will include tables for users, roles, customers, departments, managers, staff, room types, rooms, bookings, payment methods, payments, invoices, services, service usage, and room maintenance records.",
    )
    add_module(
        doc,
        "Hotel Operations Layer",
        "This layer will implement the actual hotel management features. Admins and managers can manage rooms, room types, departments, staff, bookings, invoices, and reports. Staff can assist with operational activities such as maintenance and booking status updates. Customers can register, view rooms, create bookings, request check-in/check-out, view invoices, make payments, and use hotel services.",
    )
    add_module(
        doc,
        "Containerization and Local Orchestration Layer",
        "Docker will be used to package the React frontend and FastAPI backend into containers. Docker Compose will be used to run the frontend, backend, and PostgreSQL database together in a repeatable local environment. This removes machine-specific setup issues and makes the project easy to run for team members and evaluators.",
    )
    add_module(
        doc,
        "Cloud Database and Storage Layer",
        "For production deployment, PostgreSQL will be hosted using a managed cloud database service such as Neon, Supabase, Railway PostgreSQL, or AWS RDS. Hotel-related files such as room images, customer ID proof documents, maintenance photos, or invoice PDFs will be stored using cloud object storage such as AWS S3, Cloudinary, or Supabase Storage. The database will store file URLs instead of storing large files directly.",
    )
    add_module(
        doc,
        "Cloud Deployment Layer",
        "The React frontend will be deployed as a static web application, while the FastAPI backend will be deployed as a running API service. This demonstrates the difference between frontend static hosting and backend compute hosting. The deployed frontend will communicate with the deployed backend using environment-based API configuration.",
    )
    add_module(
        doc,
        "DevOps Automation Layer",
        "GitHub Actions will be used to implement a CI/CD pipeline. On every push or pull request, the pipeline can install dependencies, validate backend code, build the React frontend, and build Docker images. On merges to the main branch, the pipeline can optionally trigger deployment to the selected cloud platforms.",
    )
    add_module(
        doc,
        "Security, Configuration, and Monitoring Layer",
        "Local development will use .env files, while production secrets such as database URLs, JWT secrets, and storage keys will be stored using GitHub Secrets or platform environment variables. Health check endpoints and basic application logging will be added to monitor backend and database status. If AWS S3 is used, IAM least-privilege credentials and signed URLs can be included for secure object access.",
    )
    add_para(
        doc,
        "By combining these layers, the project will demonstrate a complete software development lifecycle: building a real hotel management application, containerizing it, connecting it to managed cloud services, deploying it publicly, automating checks through CI/CD, and documenting the architecture for submission and demonstration.",
    )

    add_heading(doc, "Tools and Technologies")
    add_label_para(doc, "Frontend Technology: ", "React with Vite for building the hotel management dashboard and role-based user interface.")
    add_label_para(doc, "Backend Technology: ", "FastAPI with Python for REST API development, JWT authentication, business logic, and cloud service integration.")
    add_label_para(doc, "Cloud Service Provider: ", "Vercel or Netlify for frontend hosting, Render or Railway for backend hosting, and Neon or Supabase for managed PostgreSQL. AWS services such as RDS and S3 may be used if an AWS-focused implementation is preferred.")
    add_label_para(doc, "Cloud Managed Services you plan to use: ", "Managed PostgreSQL database, cloud application hosting, cloud object storage for hotel files, platform environment variables, deployment logs, and optional uptime or error monitoring.")
    add_label_para(doc, "Databases you plan to use: ", "PostgreSQL. The local version may run through Docker Compose, while the deployed version will use a managed PostgreSQL service.")
    add_label_para(doc, "Source version control: ", "Git and GitHub for repository hosting, collaboration, pull requests, and version tracking.")
    add_label_para(doc, "Tools you plan to use: ", "React, Vite, FastAPI, Python, PostgreSQL, Docker, Docker Compose, GitHub Actions, GitHub Secrets, pgAdmin/psql, Cloudinary or Supabase Storage or AWS S3, and optional Sentry or uptime monitoring.")

    add_heading(doc, "Expected Outcomes")
    add_para(
        doc,
        "The expected outcome is a complete hotel management web application with a React frontend, FastAPI backend, and PostgreSQL database. The application will provide a usable dashboard experience for multiple user roles and will support important hotel workflows such as authentication, room management, room availability checking, customer bookings, booking approval, check-in/check-out, invoice generation, payment recording, service usage, staff operations, maintenance tracking, and reports.",
    )
    add_para(
        doc,
        "From a cloud computing and DevOps perspective, the project will be considered successful if the same system can run locally using Docker Compose, connect to a managed cloud PostgreSQL database, deploy the frontend and backend to public cloud platforms, automate build validation using GitHub Actions, securely manage secrets, store files using cloud object storage, and expose health checks for monitoring. The final project should be suitable for classroom demonstration, GitHub submission, and resume presentation as a cloud-native full-stack system.",
    )

    add_heading(doc, "Tentative Timeline")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_borders(table)
    set_cell_text(table.cell(0, 0), "Week #", bold=True)
    set_cell_text(table.cell(0, 1), "Phase/Tasks", bold=True)
    timeline = [
        ("1-2", "Requirement analysis, project scope finalization, and cloud architecture planning"),
        ("3", "High-level design of frontend, backend, database, Docker setup, and deployment flow"),
        ("4", "Low-level design of API endpoints, PostgreSQL schema, UI workflows, and environment configuration"),
        ("5-6", "Implementation of hotel management modules, Dockerfiles, Docker Compose, and automated database setup"),
        ("7", "Managed cloud PostgreSQL setup and backend environment integration"),
        ("8", "Cloud storage integration for room images, documents, or invoice files"),
        ("9", "GitHub Actions CI pipeline for backend checks, frontend build, and Docker build validation"),
        ("10", "Frontend and backend cloud deployment with production environment variables"),
        ("11", "Health checks, logging, system testing, integration testing, and deployment verification"),
        ("12", "Documentation, architecture diagram, screenshots, synopsis refinement, and final presentation"),
    ]
    for week, task in timeline:
        cells = table.add_row().cells
        set_cell_text(cells[0], week)
        set_cell_text(cells[1], task)

    doc.core_properties.title = "Cloud-Native Hotel Management System Synopsis"
    doc.core_properties.subject = "Cloud Computing and DevOps mini project synopsis"
    doc.core_properties.author = "Codex"
    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
